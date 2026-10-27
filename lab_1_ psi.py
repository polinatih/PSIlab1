"""Лаб. 1: генератор экзаменационных билетов.

Зависимость: pip install openpyxl
Запуск:      python ticket_generator.py
"""
import os
import random
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook

FILE_NAME = os.path.join(os.path.dirname(os.path.abspath(__file__)), "journal.xlsx")
HEADER = ["Last name", "First name", "Номер билета", "Дата и время"]
ESC = "\x1b"

# ---------- Чтение клавиш (Windows / Linux / macOS) ----------
if os.name == "nt":
    import msvcrt

    def read_key():
        ch = msvcrt.getwch()
        if ch in ("\x00", "\xe0"):  # служебные клавиши (стрелки, F1...) — пропускаем
            msvcrt.getwch()
            return ""
        return ch
else:
    import termios
    import tty

    def read_key():
        fd = sys.stdin.fileno()
        old = termios.tcgetattr(fd)
        try:
            tty.setcbreak(fd)
            ch = sys.stdin.read(1)
        finally:
            termios.tcsetattr(fd, termios.TCSADRAIN, old)
        return ch


def read_line(prompt):
    """Читает строку посимвольно. Возвращает None, если нажат ESC."""
    print(prompt, end="", flush=True)
    buf = []
    while True:
        ch = read_key()
        if ch == ESC:
            print()
            return None
        if ch in ("\r", "\n"):
            print()
            return "".join(buf)
        if ch in ("\x08", "\x7f"):  # Backspace
            if buf:
                buf.pop()
                print("\b \b", end="", flush=True)
        elif ch == "\x03":  # Ctrl+C
            raise KeyboardInterrupt
        elif ch and ch.isprintable():
            buf.append(ch)
            print(ch, end="", flush=True)


def ask_required(prompt):
    """Повторяет запрос, пока не введена непустая строка. None — если ESC."""
    while True:
        value = read_line(prompt)
        if value is None:
            return None
        value = value.strip()
        if value:
            return value
        print("Поле не может быть пустым.")


# ---------- Работа с Excel ----------
def append_record(last_name, first_name, ticket):
    """Дописывает строку в journal.xlsx и сразу сохраняет файл."""
    if os.path.exists(FILE_NAME):
        wb = load_workbook(FILE_NAME)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(HEADER)
        for col, width in zip("ABCD", (20, 20, 15, 22)):
            ws.column_dimensions[col].width = width

    ws.append([last_name, first_name, ticket, datetime.now().strftime("%d.%m.%Y %H:%M:%S")])

    # Сохраняем во временный файл и подменяем оригинал: при сбое
    # посреди записи старый journal.xlsx остаётся целым.
    tmp_name = FILE_NAME + ".tmp"
    wb.save(tmp_name)
    try:
        os.replace(tmp_name, FILE_NAME)
    except PermissionError:
        os.remove(tmp_name)
        raise


def save_with_retry(last_name, first_name, ticket):
    """Пытается записать; если файл занят (открыт в Excel) — просит закрыть его.
    Возвращает True при успехе, False если пользователь отказался (ESC)."""
    while True:
        try:
            append_record(last_name, first_name, ticket)
            return True
        except PermissionError:
            print(
                "\nНе удалось записать в journal.xlsx: файл открыт в другой программе "
                "(скорее всего, в Excel).\n"
                "Закройте файл и нажмите Enter, чтобы повторить, или ESC, чтобы отказаться."
            )
            while True:
                key = read_key()
                if key == ESC:
                    return False
                if key in ("\r", "\n"):
                    break


# ---------- Главный цикл ----------
def main():
    print("Для выхода нажмите ESC.")
    try:
        while True:
            last_name = ask_required("Last name: ")
            if last_name is None:
                break
            first_name = ask_required("First name: ")
            if first_name is None:
                break

            ticket = random.randint(1, 20)
            print(f"Билет № {ticket}")

            if not save_with_retry(last_name, first_name, ticket):
                print("Запись не сохранена.")
            print()
    except KeyboardInterrupt:
        pass
    print("Работа завершена.")


if __name__ == "__main__":
    main()